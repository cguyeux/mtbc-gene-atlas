#!/usr/bin/env python3
"""phase68_rbtnseq_fitness.py -- P16.9 (pilote A) : couche `rbtnseq_fitness` (phénotypes conditionnels RB-TnSeq).

Ingère les phénotypes CONDITION-SPÉCIFIQUES déjà appelés par les auteurs (pas de re-dérivation de significativité)
du crible RB-TnSeq 95-conditions de Mtb (Bhatt/… PLoS Biol 2026, doi:10.1371/journal.pbio.3003529) :
- S5 « Specific phenotypes » (protéome général)  ->  résultats/phase68_rbtnseq/S_005.xlsx
- S7 « pe/ppe significant BarSeq hits »          ->  résultats/phase68_rbtnseq/S_007.xlsx

Écrit une couche `rbtnseq_fitness` sur chaque fiche présente dans l'atlas. C'est une couche de PREUVE EXPÉRIMENTALE
(phénotype de fitness conditionnel), postérieure aux fiches, distincte de MtbTnDB (phase24 ; autre compendium).

GARDE-FOUS (anti-survente, cf. scan P16) :
- un phénotype de fitness N'EST PAS une fonction prouvée -> formulé « candidat …, à valider », JAMAIS de flip de verdict.
- direction lue du SIGNE : log2/t < 0 = mutant appauvri = gène REQUIS sous la condition ; log2/t > 0 = mutant enrichi =
  la PERTE du gène confère un AVANTAGE sous la condition (signal différent, plus faible pour l'assignation de fonction).
- angle mort assumé : RB-TnSeq ne mesure pas les ~600 gènes essentiels (dont plusieurs cibles deep-dive) -> pas de signal.

Écrit EN PLACE dans site/content/genes/*.json. Lecture SEULE des xlsx.
Run: python analyses/phase68_rbtnseq_fitness.py
"""
from __future__ import annotations
import glob, json
from collections import defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
GENES = ROOT / "site" / "content" / "genes"
XDIR = ROOT / "résultats" / "phase68_rbtnseq"
DOI = "10.1371/journal.pbio.3003529"
SOURCE = f"RB-TnSeq 95-condition barcoded transposon screen, Mtb (PLoS Biol 2026, doi:{DOI})"
MAX_PH = 14                       # cap de phénotypes listés par gène (garde les |t| max)
DROP_GROUPS = {"no stress control"}


def read_table(path: Path, hdr: int) -> list[dict]:
    ws = load_workbook(path, read_only=True, data_only=True).active
    data = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c else "" for c in data[hdr]]
    return [dict(zip(header, r)) for r in data[hdr + 1:] if r[0] is not None]


def num(x):
    try:
        return round(float(x), 3)
    except (TypeError, ValueError):
        return None


def cond_label(r: dict) -> str:
    """libellé de condition robuste (S5 pH a condition=None -> utiliser 'short')."""
    c = (r.get("condition") or "").strip() if r.get("condition") else ""
    if c and c.lower() != "none":
        return c
    return (str(r.get("short") or "").strip()) or "(unspecified)"


def collect(rows: list[dict], is_s7: bool) -> dict[str, list[dict]]:
    """rv -> liste de phénotypes {condition, group, log2, t, direction, detail} dédupliqués par condition (|t| max)."""
    by_gene: dict[str, dict[str, dict]] = defaultdict(dict)
    for r in rows:
        rv = (r.get("locusId") or "").strip()
        if not rv:
            continue
        group = (r.get("group") or ("pe/ppe" if is_s7 else "")).strip()
        if group in DROP_GROUPS:
            continue
        log2, t = num(r.get("log2")), num(r.get("t"))
        if log2 is None or t is None:
            continue
        cond = cond_label(r)
        ph = {
            "condition": cond,
            "group": group or None,
            "log2": log2,
            "t": t,
            "direction": "required" if log2 < 0 else "advantage",  # <0 : mutant appauvri = gène requis
            "detail": (str(r.get("short")) if r.get("short") else "") or None,
        }
        prev = by_gene[rv].get(cond)
        if prev is None or abs(t) > abs(prev["t"]):
            by_gene[rv][cond] = ph
    return {rv: sorted(d.values(), key=lambda p: -abs(p["t"])) for rv, d in by_gene.items()}


def make_lead(rv: str, phenos: list[dict], verdict: str) -> str:
    """Texte de lead honnête, gradué par direction ; cadrage 'candidat / à valider' surtout pour les dark."""
    top = phenos[0]
    cond, grp, d = top["condition"], (top["group"] or "condition"), top["direction"]
    gmap = {"carbon source": "carbon-source", "nitrogen source": "nitrogen-source",
            "pH": "pH", "stress": "stress", "pe/ppe": "stress/antibiotic"}
    gl = gmap.get(grp, grp)
    if d == "required":
        body = (f"transposon mutants are specifically depleted under {cond} ({gl}) "
                f"(RB-TnSeq log2={top['log2']}, t={top['t']}), i.e. the gene is required for fitness in that condition")
        hint = f" — candidate involvement in the {cond}-related process/utilisation"
    else:
        body = (f"transposon mutants are specifically enriched under {cond} ({gl}) "
                f"(RB-TnSeq log2={top['log2']}, t={top['t']}), i.e. loss of the gene confers a fitness advantage there")
        hint = f" — candidate {cond}-susceptibility/activation factor whose disruption is beneficial"
    more = f"; {len(phenos)} specific phenotype(s) total" if len(phenos) > 1 else ""
    tail = ", to validate (a fitness phenotype is not a proven function)" if verdict == "dark" else ""
    return body + hint + more + tail + "."


def main() -> None:
    s5 = read_table(XDIR / "S_005.xlsx", 12)
    s7 = read_table(XDIR / "S_007.xlsx", 7)
    g5 = collect(s5, is_s7=False)
    g7 = collect(s7, is_s7=True)
    pe_ppe = set(g7)
    # fusion : un gène pe/ppe peut n'apparaître que dans S7
    all_genes = set(g5) | set(g7)

    verd = {}
    for f in glob.glob(str(GENES / "*.json")):
        verd[Path(f).stem] = json.loads(Path(f).read_text()).get("verdict")

    written = 0
    dark_leads = []
    grp_counter = defaultdict(int)
    for f in glob.glob(str(GENES / "*.json")):
        d = json.loads(Path(f).read_text())
        rv = d["rv"]
        if rv not in all_genes:
            continue
        phenos = (g5.get(rv, []) + g7.get(rv, []))
        # dédup finale par condition (S5/S7 peuvent recouvrir), |t| max
        seen = {}
        for p in sorted(phenos, key=lambda p: -abs(p["t"])):
            seen.setdefault(p["condition"], p)
        phenos = sorted(seen.values(), key=lambda p: -abs(p["t"]))
        for p in phenos:
            grp_counter[p["group"]] += 1
        verdict = d.get("verdict")
        layer = {
            "n_phenotypes": len(phenos),
            "groups": sorted({p["group"] for p in phenos if p["group"]}),
            "phenotypes": phenos[:MAX_PH],
            "is_pe_ppe": rv in pe_ppe,
            "lead": make_lead(rv, phenos, verdict),
            "source": SOURCE,
        }
        d["rbtnseq_fitness"] = layer
        Path(f).write_text(json.dumps(d, indent=2, ensure_ascii=False))
        written += 1
        if verdict == "dark":
            dark_leads.append((rv, layer["lead"]))

    print(f"couche `rbtnseq_fitness` écrite sur {written} fiches présentes dans l'atlas.")
    print(f"  (S5 gènes {len(g5)}, S7 pe/ppe {len(g7)}, union {len(all_genes)})")
    print(f"  phénotypes par groupe : {dict(grp_counter)}")
    print(f"\n>>> {len(dark_leads)} gène(s) DARK reçoivent un lead de contexte RB-TnSeq :")
    for rv, lead in dark_leads:
        print(f"  {rv}: {lead}")
    print("\nRappel honnêteté : angle mort essentiels (non mesurés) ; phénotype = lead, jamais flip.")


if __name__ == "__main__":
    main()
