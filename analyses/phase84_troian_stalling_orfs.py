#!/usr/bin/env python3
"""phase84 — Troian 2026 VapC4 ribosome-stalling ORFs -> microproteome track (P16.5b).

Source: Troian EA, Barth VC, Chauhan U, ..., Husson RN, Woychik NA. "Harnessing
toxin-mediated ribosome stalling as a complementary tool to annotate the
Mycobacterium tuberculosis genome." Nucleic Acids Res. 2026;54(7):gkag252
(doi:10.1093/nar/gkag252; PMC13096801). Supplementary hand-off by CG (OUP anti-bot):
résultats/phase84_troian/extracted/Supplemental Tables 1-5.xlsx.

Method is ORTHOGONAL to de Souza (phase83): the VapC4 toxin cleaves tRNA-Cys, so
ribosomes stall at Cys codons; mapping the stalled-ribosome footprints (5'-OH
RNA-seq) detects unannotated ORFs. Table 1 lists 96 unannotated ORFs; existence
evidence is the stalling footprint, with two orthogonal cross-validations recorded
per ORF: MS (ProteomeXchange datasets, Supplemental Table 2) and Ribo-RET (Smith
et al. 2022, "Found in Smith" column).

Measured (analyses run before ingest): 0/96 overlap a de Souza microprotein on the
same strand -> the two methods are non-redundant; all 96 are additive. 5 ORFs now
carry an official Rv tag (Rv0815A/Rv0485A/Rv2742A/Rv2391A/Rv2334A) yet are ABSENT
from this atlas's 3906 CDS -> flagged as atlas gaps.

Writes content/microproteins_troian.json (ingest globs content/microproteins*.json).
Read-only on the canonical gene set. Overlap class assigned by GEOMETRY vs H37Rv CDS.
"""
from __future__ import annotations
import json, os, re, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "résultats", "phase84_troian", "extracted", "Supplemental Tables 1-5.xlsx")
GFF = os.path.join(ROOT, "..", "investigate_phylo", "resources", "NC_000962.3.gff3")
OUT = os.path.join(ROOT, "site", "content", "microproteins_troian.json")
CITATION = ("Troian EA et al., Nucleic Acids Res. 2026;54(7):gkag252 "
            "(doi:10.1093/nar/gkag252; PMC13096801); VapC4 ribosome-stalling proteogenomics")
DETECTION = "VapC4 ribosome-stalling (5'-OH RNA-seq)"


def load_cds():
    cds = []
    for line in open(GFF):
        if line.startswith("#"):
            continue
        p = line.rstrip("\n").split("\t")
        if len(p) < 9 or p[2] != "CDS":
            continue
        m = re.search(r"locus_tag=([^;]+)", p[8])
        cds.append((int(p[3]), int(p[4]), p[6], m.group(1) if m else "?"))
    return cds


def classify(a, b, strand, cds):
    hits = [c for c in cds if not (c[1] < a or c[0] > b)]
    if not hits:
        return "intergenic", None
    same = [c for c in hits if c[2] == strand]
    return ("sense_overlap", same[0][3]) if same else ("antisense", hits[0][3])


def sheet(name):
    ws = load_workbook(XLSX, data_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows)
              if r and any(str(c).startswith(("Unannotated ORF", "Dataset")) for c in r if c))
    hdr = [str(c).strip() if c is not None else "" for c in rows[hi]]
    return hdr, rows[hi + 1:]


def main():
    # MS validation: ORF -> set(PXD datasets) from Supplemental Table 2
    ms = {}
    try:
        h2, r2 = sheet("Supplemental Table 2")
        di = h2.index("Dataset") if "Dataset" in h2 else 0
        oi = h2.index("Unannotated ORF") if "Unannotated ORF" in h2 else 1
        for r in r2:
            if not r or not r[oi]:
                continue
            ms.setdefault(str(r[oi]).strip(), set()).add(str(r[di]).strip())
    except Exception as e:
        print(f"  [warn] Table 2 (MS) illisible: {e}")

    cds = load_cds()
    h, rows = sheet("Supplemental Table 1")
    ix = {k: h.index(k) for k in h}
    items = []
    for r in rows:
        if not r or not r[ix["Unannotated ORF"]]:
            continue
        oid = str(r[ix["Unannotated ORF"]]).strip()
        orient = str(r[ix["Orientation"]]).strip() or "+"
        a, b = int(r[ix["Start"]]), int(r[ix["Stop"]])
        start, end = min(a, b), max(a, b)
        strand = "+" if orient.startswith("+") else "-"
        aa = str(r[ix["Amino Acid Sequence"]] or "").rstrip("*").strip()
        oc, loc = classify(start, end, strand, cds)
        annotated = r[ix["Annotated?"]]
        annotated = str(annotated).strip() if annotated not in (None, "", "No") else None
        smith = str(r[ix["Found in Smith et al. 2022"]]).strip().lower() == "yes"
        notes = str(r[ix["Notes"]] or "").strip()
        transcript_class = notes if notes in ("Leaderless", "Leadered") else None
        ms_sets = sorted(ms.get(oid, []))
        # per-ORF evidence string, transparent about what supports it
        ev = [DETECTION]
        if ms_sets:
            ev.append(f"MS-validated (ProteomeXchange {', '.join(ms_sets)})")
        if smith:
            ev.append("cross-validated by Ribo-RET (Smith et al. 2022)")
        items.append({
            "id": oid,
            "kind": "ribosome-stalling ORF",
            "overlap_class": oc, "overlap_locus": loc,
            "start": start, "end": end, "strand": strand,
            "length_aa": len(aa), "sequence": aa or None,
            "ms_evidence": "; ".join(ev),
            "detection_method": DETECTION,
            "source": CITATION,
            "extra": {
                "n_cysteines": r[ix["# Cysteines"]],
                "transcript_class": transcript_class,
                "ms_validated": bool(ms_sets),
                "ms_datasets": ms_sets or None,
                "ribo_ret_smith2022": smith,
                "now_annotated_as": annotated,
                "note": (f"now officially annotated as {annotated}, but ABSENT from this atlas's "
                         f"canonical set (added to H37Rv after build)" if annotated else None),
            },
        })

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(items, open(OUT, "w"), ensure_ascii=False, indent=2)

    from collections import Counter
    oc = Counter(it["overlap_class"] for it in items)
    n_ms = sum(1 for it in items if it["extra"]["ms_validated"])
    n_smith = sum(1 for it in items if it["extra"]["ribo_ret_smith2022"])
    n_small = sum(1 for it in items if it["length_aa"] <= 50)
    n_ann = sum(1 for it in items if it["extra"]["now_annotated_as"])
    print(f"phase84: {len(items)} ORF stalling VapC4 -> {OUT}")
    print(f"  classification géométrie H37Rv: {dict(oc)}")
    print(f"  MS-validés: {n_ms} | recoupés Ribo-RET/Smith2022: {n_smith} | "
          f"smORF <=50 aa: {n_small} | déjà annotés (absents atlas): {n_ann}")
    print(f"  gaps atlas (Rv officiels manquants): "
          f"{[it['extra']['now_annotated_as'] for it in items if it['extra']['now_annotated_as']]}")


if __name__ == "__main__":
    main()
