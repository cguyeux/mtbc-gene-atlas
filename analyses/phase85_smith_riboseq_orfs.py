#!/usr/bin/env python3
"""phase85 — Smith 2022 Ribo-seq/Ribo-RET novel ORFs -> microproteome track (P16.5c, light entry).

Source: Smith C, Canestrari JG, Wang AJ et al. "Pervasive translation in
Mycobacterium tuberculosis." eLife 2022;11:e73980 (doi:10.7554/eLife.73980;
PMC9094748). OPEN ACCESS. Supplementary File 1 fetched from the eLife CDN
(résultats/phase85_smith/supp1.xlsx, sheet "Supplementary File 1A").

/challenge verdict (P16.5c): the piste's "de novo Ribo-seq reanalysis" is NOT done
— it would rediscover this peer-reviewed catalogue and risk our own ORF-calling
errors. Instead we ingest Smith's published catalogue as a THIRD orthogonal source
(ribosome-profiling translation evidence), completing the MS / VapC4-stalling /
Ribo-seq triad. The true de novo reanalysis stays deferred.

Scope (chosen with CG): the CONFIDENT subset only. Table 1A classifies ORFs as
Annotated / Isoform / Novel; we take "Novel" that are (a) NOT already in the track
(no same-strand coordinate overlap with the 168 de Souza+Troian entries) AND (b)
carry independent evidence of being a real coding unit — MS-detected (>=1 peptide
or an MS FDR) OR a significant protein-coding G/C-skew signature (p<0.05). This
excludes the ~323 ribosome-occupancy-only calls (median 26 aa) of "pervasive
translation", which are likely regulatory/spurious and would oversell the track.

IDs are coordinate-based (Smith_<start>_<strand>) since 1A has no ORF id. Overlap
class assigned by GEOMETRY vs H37Rv CDS. Writes content/microproteins_smith.json
(ingest globs content/microproteins*.json). Read-only on the canonical gene set.
"""
from __future__ import annotations
import json, os, re, warnings
warnings.filterwarnings("ignore")
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "résultats", "phase85_smith", "supp1.xlsx")
GFF = os.path.join(ROOT, "..", "investigate_phylo", "resources", "NC_000962.3.gff3")
OUT = os.path.join(ROOT, "site", "content", "microproteins_smith.json")
EXISTING = ["site/content/microproteins.json", "site/content/microproteins_troian.json"]
CITATION = ("Smith C, Canestrari JG, Wang AJ et al., eLife 2022;11:e73980 "
            "(doi:10.7554/eLife.73980; PMC9094748); ribosome-profiling translatome")
DETECTION = "Ribo-seq / Ribo-RET translation (Smith 2022)"


def fnum(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def load_cds():
    cds = []
    for line in open(GFF):
        if line.startswith("#"):
            continue
        p = line.rstrip("\n").split("\t")
        if len(p) < 9 or p[2] != "CDS":
            continue
        m = re.search(r"locus_tag=([^;]+)", p[8])
        cds.append((int(p[3]), int(p[4]), p[6], m.group(1) if m else "?"))
    return cds


def classify(a, b, strand, cds):
    hits = [c for c in cds if not (c[1] < a or c[0] > b)]
    if not hits:
        return "intergenic", None
    same = [c for c in hits if c[2] == strand]
    return ("sense_overlap", same[0][3]) if same else ("antisense", hits[0][3])


def load_existing():
    ex = []
    for f in EXISTING:
        fp = os.path.join(ROOT, f)
        if os.path.exists(fp):
            for m in json.load(open(fp)):
                if m.get("start"):
                    ex.append((min(m["start"], m["end"]), max(m["start"], m["end"]), m["strand"]))
    return ex


def main():
    ws = load_workbook(XLSX, data_only=True)["Supplementary File 1A"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and any(str(c).strip() == "Strand" for c in r if c))
    h = [str(c).strip() if c is not None else "" for c in rows[hi]]
    ix = {x: i for i, x in enumerate(h)}

    def col(prefix: str) -> int:
        i = next((i for k, i in ix.items() if k.startswith(prefix)), None)
        if i is None:
            raise KeyError(f"colonne '{prefix}...' absente de Supplementary File 1A")
        return i

    c_cls = ix["Classification"]; c_str = ix["Strand"]
    c_tss = col("TSS/Start"); c_stop = col("Stop Codon Coord"); c_prot = col("Protein Sequence")
    c_pep = col("# MS Peptides"); c_fdr = col("MS FDR"); c_dens = col("Relative Ribosome")
    c_gc = col("G/C Skew"); c_ov = col("Annotated Gene Overlap")

    cds = load_cds()
    existing = load_existing()

    def is_new(a, b, strand):
        return not any(not (b < ea or a > eb) and strand == es for ea, eb, es in existing)

    items, seen = [], set()
    n_ms = n_gc = 0
    for r in rows[hi + 1:]:
        if not r or str(r[c_cls]) != "Novel":
            continue
        a, b = fnum(r[c_tss]), fnum(r[c_stop])
        if a is None or b is None:
            continue
        a, b = int(a), int(b)
        start, end = min(a, b), max(a, b)
        strand = "+" if str(r[c_str]).strip().startswith("+") else "-"
        if not is_new(start, end, strand):
            continue
        pep = fnum(r[c_pep])
        has_ms = (pep is not None and pep >= 1) or bool(r[c_fdr] and str(r[c_fdr]).strip() not in ("", "None", "NA", "-"))
        gcp = fnum(r[c_gc])
        gc_sig = gcp is not None and gcp < 0.05
        if not (has_ms or gc_sig):            # CONFIDENT subset only
            continue
        aa = str(r[c_prot] or "").rstrip("*").strip()
        oc, loc = classify(start, end, strand, cds)
        sid = f"Smith_{start}_{'p' if strand == '+' else 'm'}"
        if sid in seen:
            continue
        seen.add(sid)
        dens = fnum(r[c_dens])
        ev = [DETECTION.replace(" (Smith 2022)", "")]
        if dens is not None:
            ev[0] += f" (relative ribosome density {dens:.2f})"
        if has_ms:
            ev.append(f"MS-detected ({int(pep)} peptide(s))" if pep else "MS-detected")
            n_ms += 1
        if gc_sig:
            ev.append(f"protein-coding G/C-skew signature (p={gcp:g})")
            n_gc += 1
        items.append({
            "id": sid,
            "kind": "ribosome-profiling ORF",
            "overlap_class": oc, "overlap_locus": loc,
            "start": start, "end": end, "strand": strand,
            "length_aa": len(aa), "sequence": aa or None,
            "ms_evidence": "; ".join(ev),
            "detection_method": DETECTION,
            "source": CITATION,
            "extra": {
                "classification": "Novel (unannotated translated ORF)",
                "relative_ribosome_density": dens,
                "ms_peptides": int(pep) if pep else 0,
                "gc_skew_pvalue": gcp,
                "coding_signature": bool(gc_sig),
                "annotated_gene_overlap": (str(r[c_ov]).strip() if c_ov is not None and r[c_ov] else None),
            },
        })

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(items, open(OUT, "w"), ensure_ascii=False, indent=2)
    from collections import Counter
    oc = Counter(it["overlap_class"] for it in items)
    print(f"phase85: {len(items)} ORF Ribo-seq Smith (sous-ensemble confiant) -> {OUT}")
    print(f"  support: MS={n_ms}, signature codante G/C={n_gc} (union, un ORF peut avoir les deux)")
    print(f"  classification géométrie: {dict(oc)}")
    lens = sorted(it["length_aa"] for it in items)
    if lens:
        print(f"  longueur aa: min={lens[0]} med={lens[len(lens)//2]} max={lens[-1]}")


if __name__ == "__main__":
    main()
